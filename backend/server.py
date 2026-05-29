"""
FerrePro — Flask API Server
Connects to Supabase/Postgres (or Turso legacy fallback) and provides REST endpoints.
Deployable on Vercel as serverless functions.
"""
import os
import json
import base64
import signal
from datetime import datetime, date
from flask import Flask, request, jsonify, send_file, send_from_directory, render_template
from flask_cors import CORS
from io import BytesIO
import sys, uuid, threading

# Get the project root directory (ferrepro folder)
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
sys.path.insert(0, PROJECT_ROOT)

import backend.db as db

app = Flask(__name__, 
    static_url_path='', 
    static_folder=os.path.join(PROJECT_ROOT, 'static'),
    template_folder=os.path.join(PROJECT_ROOT, 'templates'))
CORS(app)

# ─── Auth (simple) ───
ADMIN_PASS = os.getenv("ADMIN_PASS", "ferrepro2026")


def require_auth():
    """Check Authorization header for admin password."""
    auth = request.headers.get("Authorization", "")
    if not auth.startswith("Bearer "):
        return None
    token = auth.replace("Bearer ", "")
    return token == ADMIN_PASS


# ═══════════════════════════════════════
#  AUTH
# ═══════════════════════════════════════

@app.route("/api/login", methods=["POST"])
def login():
    data = request.get_json(silent=True) or {}
    user = data.get("username", "")
    password = data.get("password", "")
    if user == "admin" and password == ADMIN_PASS:
        return jsonify({"ok": True, "token": ADMIN_PASS})
    return jsonify({"ok": False}), 401


# ═══════════════════════════════════════
#  ARTÍCULOS
# ═══════════════════════════════════════

@app.route("/api/articulos", methods=["GET"])
def get_articulos():
    search = request.args.get("search", "")
    dept = request.args.get("departamento", "")
    stock_filter = request.args.get("stock", "")
    
    sql = """SELECT a.*, p.empresa as proveedor_nombre 
             FROM articulos a LEFT JOIN proveedores p ON a.proveedor_id=p.id"""
    conds = []
    params = []
    
    if search:
        conds.append("(LOWER(a.nombre) LIKE ? OR LOWER(a.codigo) LIKE ?)")
        params.extend([f"%{search.lower()}%", f"%{search.lower()}%"])
    if dept:
        conds.append("a.departamento = ?")
        params.append(dept)
    if stock_filter == "bajo":
        conds.append("a.stock <= a.stock_min")
    elif stock_filter == "ok":
        conds.append("a.stock > a.stock_min")
    
    if conds:
        sql += " WHERE " + " AND ".join(conds)
    sql += " ORDER BY a.nombre"
    
    try:
        items = db.query(sql, params) if params else db.query(sql)
        return jsonify(items)
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/articulos/<int:art_id>", methods=["GET"])
def get_articulo(art_id):
    items = db.query("SELECT * FROM articulos WHERE id=?", [art_id])
    if not items:
        return jsonify({"error": "No encontrado"}), 404
    return jsonify(items[0])


@app.route("/api/articulos/buscar-por-barras", methods=["GET"])
def get_articulo_by_barcode():
    codigo = request.args.get("codigo", "")
    if not codigo:
        return jsonify({"error": "Código requerido"}), 400
    items = db.query("SELECT * FROM articulos WHERE codigo_barras=?", [codigo])
    if not items:
        return jsonify({"error": "Artículo no encontrado"}), 404
    return jsonify(items[0])


@app.route("/api/articulos/<int:art_id>/barcode", methods=["POST"])
def update_articulo_barcode(art_id):
    data = request.get_json(silent=True) or {}
    codigo_barras = data.get("codigo_barras", "")
    db.execute("UPDATE articulos SET codigo_barras=? WHERE id=?", [codigo_barras, art_id])
    return jsonify({"ok": True})


@app.route("/api/articulos", methods=["POST"])
def create_articulo():
    data = request.get_json(silent=True) or {}
    # Validate required
    if not data.get("nombre"):
        return jsonify({"error": "El nombre es requerido"}), 400
    
    sql = """INSERT INTO articulos 
             (codigo, nombre, departamento, precio_costo, precio_venta, stock, stock_min, descripcion, proveedor_id, unidad)
             VALUES (?,?,?,?,?,?,?,?,?,?)"""
    try:
        db.execute(sql, [
            data.get("codigo", ""),
            data["nombre"],
            data.get("departamento", "Ferretería"),
            float(data.get("precio_costo", 0)),
            float(data.get("precio_venta", 0)),
            int(data.get("stock", 0)),
            int(data.get("stock_min", 5)),
            data.get("descripcion", ""),
            data.get("proveedor_id"),
            data.get("unidad", "u.")
        ])
        return jsonify({"ok": True})
    except Exception as e:
        if "UNIQUE" in str(e):
            return jsonify({"error": "El código ya existe"}), 409
        return jsonify({"error": str(e)}), 500


@app.route("/api/articulos/batch", methods=["POST"])
def create_articulos_batch():
    data = request.get_json(silent=True) or []
    if not isinstance(data, list) or not data:
        return jsonify({"error": "Se requiere una lista de articulos"}), 400
    sql_parts = []
    params = []
    for item in data:
        # minimal validation
        nombre = item.get("nombre")
        if not nombre:
            return jsonify({"error": "Cada articulo necesita 'nombre'"}), 400
        sql_parts.append("INSERT INTO articulos (codigo, nombre, departamento, precio_costo, precio_venta, stock, stock_min, descripcion, proveedor_id, unidad) VALUES (?,?,?,?,?,?,?,?,?,?);")
        params.extend([
            item.get("codigo", ""),
            nombre,
            item.get("departamento", "Ferretería"),
            float(item.get("precio_costo", 0)),
            float(item.get("precio_venta", 0)),
            int(item.get("stock", 0)),
            int(item.get("stock_min", 5)),
            item.get("descripcion", ""),
            item.get("proveedor_id"),
            item.get("unidad", "u.")
        ])
    full_sql = "BEGIN TRANSACTION; " + " ".join(sql_parts) + " COMMIT;"
    try:
        db.execute(full_sql, params)
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/articulos/<int:art_id>", methods=["PUT"])
def update_articulo(art_id):
    data = request.get_json(silent=True) or {}
    if not data.get("nombre"):
        return jsonify({"error": "El nombre es requerido"}), 400
    
    sql = """UPDATE articulos SET codigo=?, nombre=?, departamento=?, precio_costo=?, precio_venta=?,
             stock=?, stock_min=?, descripcion=?, proveedor_id=?, unidad=? WHERE id=?"""
    try:
        db.execute(sql, [
            data.get("codigo", ""),
            data["nombre"],
            data.get("departamento", "Ferretería"),
            float(data.get("precio_costo", 0)),
            float(data.get("precio_venta", 0)),
            int(data.get("stock", 0)),
            int(data.get("stock_min", 5)),
            data.get("descripcion", ""),
            data.get("proveedor_id"),
            data.get("unidad", "u."),
            art_id
        ])
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/articulos/<int:art_id>", methods=["DELETE"])
def delete_articulo(art_id):
    db.execute("DELETE FROM articulos WHERE id=?", [art_id])
    return jsonify({"ok": True})


@app.route("/api/articulos/limpiar-duplicados", methods=["POST"])
def limpiar_duplicados_articulos():
    """
    Identifica y elimina artículos duplicados.
    Prioriza 'codigo': si existe, se usa como identificador único.
    Si 'codigo' está vacío, usa la combinación 'nombre' y 'departamento'.
    Conserva la primera entrada encontrada para cada duplicado y elimina las subsiguientes.
    """
    try:
        # Obtener todos los artículos con los campos necesarios para identificar duplicados
        all_articles = db.query("SELECT id, codigo, nombre, departamento FROM articulos")
        
        duplicates_to_delete_ids = []
        
        # Diccionarios para rastrear elementos vistos
        # {codigo: id_del_articulo_a_conservar}
        seen_codes = {}
        # {(nombre, departamento): id_del_articulo_a_conservar}
        seen_names_depts = {}

        for article in all_articles:
            current_id = article['id']
            codigo = (article['codigo'] or '').strip().lower() # Normalizar y manejar códigos nulos/vacíos
            nombre = (article['nombre'] or '').strip().lower()
            departamento = (article['departamento'] or '').strip().lower()

            if codigo: # Si el artículo tiene un código, usarlo como identificador principal
                if codigo in seen_codes:
                    duplicates_to_delete_ids.append(current_id)
                else:
                    seen_codes[codigo] = current_id
            else: # Si el código está vacío, usar nombre y departamento
                key = (nombre, departamento)
                if key in seen_names_depts:
                    duplicates_to_delete_ids.append(current_id)
                else:
                    seen_names_depts[key] = current_id
        
        if duplicates_to_delete_ids:
            placeholders = ', '.join(['?' for _ in duplicates_to_delete_ids])
            db.execute(f"DELETE FROM articulos WHERE id IN ({placeholders})", duplicates_to_delete_ids)
            
        return jsonify({"ok": True, "deleted_count": len(duplicates_to_delete_ids)})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ═══════════════════════════════════════
#  PROVEEDORES
# ═══════════════════════════════════════

@app.route("/api/proveedores", methods=["GET"])
def get_proveedores():
    search = request.args.get("search", "")
    sql = """SELECT p.*, (SELECT COUNT(*) FROM facturas f WHERE f.proveedor_id=p.id) as num_fact
             FROM proveedores p"""
    if search:
        sql += " WHERE LOWER(p.empresa) LIKE ? OR LOWER(p.contacto) LIKE ?"
        s = f"%{search.lower()}%"
        items = db.query(sql, [s, s])
    else:
        items = db.query(sql)
    return jsonify(items)


@app.route("/api/proveedores/<int:prov_id>", methods=["GET"])
def get_proveedor(prov_id):
    items = db.query("SELECT * FROM proveedores WHERE id=?", [prov_id])
    if not items:
        return jsonify({"error": "No encontrado"}), 404
    return jsonify(items[0])


@app.route("/api/proveedores", methods=["POST"])
def create_proveedor():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"error": "JSON inválido"}), 400

    if not data.get("empresa"):
        return jsonify({"error": "El nombre de la empresa es requerido"}), 400
    
    sql = """INSERT INTO proveedores 
             (empresa, contacto, telefono, email, departamento, direccion, rnc, dias_credito, notas)
             VALUES (?,?,?,?,?,?,?,?,?)"""
    db.execute(sql, [
        data["empresa"],
        data.get("contacto", ""),
        data.get("telefono", ""),
        data.get("email", ""),
        data.get("departamento", ""),
        data.get("direccion", ""),
        data.get("rnc", ""),
        int(data.get("dias_credito", 30)),
        data.get("notas", "")
    ])
    return jsonify({"ok": True})


@app.route("/api/proveedores/batch", methods=["POST"])
def create_proveedores_batch():
    data = request.get_json(silent=True) or []
    if not isinstance(data, list) or not data:
        return jsonify({"error": "Se requiere una lista de proveedores"}), 400
    sql_parts = []
    params = []
    for item in data:
        empresa = item.get("empresa")
        if not empresa:
            return jsonify({"error": "Cada proveedor necesita 'empresa'"}), 400
        sql_parts.append("INSERT INTO proveedores (empresa, contacto, telefono, email, departamento, direccion, rnc, dias_credito, notas) VALUES (?,?,?,?,?,?,?,?,?);")
        params.extend([
            empresa,
            item.get("contacto", ""),
            item.get("telefono", ""),
            item.get("email", ""),
            item.get("departamento", ""),
            item.get("direccion", ""),
            item.get("rnc", ""),
            int(item.get("dias_credito", 30)),
            item.get("notas", "")
        ])
    full_sql = "BEGIN TRANSACTION; " + " ".join(sql_parts) + " COMMIT;"
    try:
        db.execute(full_sql, params)
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/proveedores/<int:prov_id>", methods=["PUT"])
def update_proveedor(prov_id):
    data = request.get_json(silent=True) or {}
    if not data.get("empresa"):
        return jsonify({"error": "El nombre de la empresa es requerido"}), 400
    
    sql = """UPDATE proveedores SET empresa=?, contacto=?, telefono=?, email=?, departamento=?,
             direccion=?, rnc=?, dias_credito=?, notas=? WHERE id=?"""
    db.execute(sql, [
        data["empresa"],
        data.get("contacto", ""),
        data.get("telefono", ""),
        data.get("email", ""),
        data.get("departamento", ""),
        data.get("direccion", ""),
        data.get("rnc", ""),
        int(data.get("dias_credito", 30)),
        data.get("notas", ""),
        prov_id
    ])
    return jsonify({"ok": True})


@app.route("/api/proveedores/<int:prov_id>", methods=["DELETE"])
def delete_proveedor(prov_id):
    # First, set proveedor_id to NULL in articulos referencing this proveedor
    db.execute("UPDATE articulos SET proveedor_id=NULL WHERE proveedor_id=?", [prov_id])
    db.execute("DELETE FROM proveedores WHERE id=?", [prov_id])
    return jsonify({"ok": True})


# ═══════════════════════════════════════
#  FACTURAS
# ═══════════════════════════════════════

@app.route("/api/facturas", methods=["GET"])
def get_facturas():
    search = request.args.get("search", "")
    estado = request.args.get("estado", "")
    fecha_inicio = request.args.get("fecha_inicio", "")
    fecha_fin = request.args.get("fecha_fin", "")
    proveedor_id = request.args.get("proveedor_id", "")
    sql = """SELECT f.*, p.empresa FROM facturas f 
             LEFT JOIN proveedores p ON f.proveedor_id=p.id"""
    conds = []
    params = []
    
    if search:
        conds.append("(LOWER(f.numero) LIKE ? OR LOWER(p.empresa) LIKE ?)")
        params.extend([f"%{search.lower()}%", f"%{search.lower()}%"])
    if estado:
        conds.append("f.estado = ?")
        params.append(estado)
    if fecha_inicio:
        conds.append("f.fecha_emision >= ?")
        params.append(fecha_inicio)
    if fecha_fin:
        conds.append("f.fecha_emision <= ?")
        params.append(fecha_fin)
    if proveedor_id:
        conds.append("f.proveedor_id = ?")
        params.append(proveedor_id)
    
    if conds:
        sql += " WHERE " + " AND ".join(conds)
    sql += " ORDER BY f.fecha_vencimiento ASC"
    
    items = db.query(sql, params) if params else db.query(sql)
    return jsonify(items)


@app.route("/api/facturas/<int:fact_id>", methods=["GET"])
def get_factura(fact_id):
    items = db.query("SELECT f.*, p.empresa FROM facturas f LEFT JOIN proveedores p ON f.proveedor_id=p.id WHERE f.id=?", [fact_id])
    if not items:
        return jsonify({"error": "No encontrado"}), 404
    return jsonify(items[0])


@app.route("/api/facturas", methods=["POST"])
def create_factura():
    data = request.get_json(silent=True) or {}
    if not data.get("numero") or not data.get("fecha_emision") or not data.get("fecha_vencimiento"):
        return jsonify({"error": "Número, emisión y vencimiento son requeridos"}), 400
    
    sql = """INSERT INTO facturas 
             (numero, proveedor_id, monto, fecha_emision, fecha_vencimiento, estado, descripcion)
             VALUES (?,?,?,?,?,?,?)"""
    db.execute(sql, [
        data["numero"],
        data.get("proveedor_id"),
        float(data.get("monto", 0)),
        data["fecha_emision"],
        data["fecha_vencimiento"],
        data.get("estado", "pendiente"),
        data.get("descripcion", "")
    ])
    return jsonify({"ok": True})


@app.route("/api/facturas/batch", methods=["POST"])
def create_facturas_batch():
    data = request.get_json(silent=True) or []
    if not isinstance(data, list) or not data:
        return jsonify({"error": "Se requiere una lista de facturas"}), 400
    sql_parts = []
    params = []
    for item in data:
        numero = item.get("numero")
        fecha_emision = item.get("fecha_emision")
        fecha_vencimiento = item.get("fecha_vencimiento")
        if not numero or not fecha_emision or not fecha_vencimiento:
            return jsonify({"error": "Cada factura requiere numero, fecha_emision y fecha_vencimiento"}), 400
        sql_parts.append("INSERT INTO facturas (numero, proveedor_id, monto, fecha_emision, fecha_vencimiento, estado, descripcion) VALUES (?,?,?,?,?,?,?);")
        params.extend([
            numero,
            item.get("proveedor_id"),
            float(item.get("monto", 0)),
            fecha_emision,
            fecha_vencimiento,
            item.get("estado", "pendiente"),
            item.get("descripcion", "")
        ])
    full_sql = "BEGIN TRANSACTION; " + " ".join(sql_parts) + " COMMIT;"
    try:
        db.execute(full_sql, params)
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/facturas/<int:fact_id>", methods=["PUT"])
def update_factura(fact_id):
    data = request.get_json(silent=True) or {}
    sql = """UPDATE facturas SET numero=?, proveedor_id=?, monto=?, fecha_emision=?, 
             fecha_vencimiento=?, estado=?, descripcion=? WHERE id=?"""
    db.execute(sql, [
        data.get("numero", ""),
        data.get("proveedor_id"),
        float(data.get("monto", 0)),
        data.get("fecha_emision", ""),
        data.get("fecha_vencimiento", ""),
        data.get("estado", "pendiente"),
        data.get("descripcion", ""),
        fact_id
    ])
    return jsonify({"ok": True})


@app.route("/api/facturas/<int:fact_id>/pagar", methods=["POST"])
def pagar_factura(fact_id):
    db.execute("UPDATE facturas SET estado='pagada', fecha_pago=? WHERE id=?", [datetime.now().isoformat(), fact_id])
    return jsonify({"ok": True})


@app.route("/api/lista-compras", methods=["GET"])
def get_lista_compras():
    departamento = request.args.get("departamento", "")
    filtro = request.args.get("filtro", "")
    sql = "SELECT id, nombre, COALESCE(cantidad, 1) AS cantidad, articulo_id, departamento, COALESCE(estado, 'pendiente') AS estado, COALESCE(encargado, '') AS encargado, creado_at, completado_at FROM lista_compras"
    params = []
    conds = []
    if filtro == "historial":
        conds.append("estado IN ('recibido', 'faltante')")
    elif not filtro:
        conds.append("estado = 'pendiente'")
    
    if departamento:
        conds.append("departamento = ?")
        params.append(departamento)
    
    if conds:
        sql += " WHERE " + " AND ".join(conds)
    sql += " ORDER BY id DESC"
    items = db.query(sql, params if params else None)
    return jsonify(items)


@app.route("/api/lista-compras/<int:item_id>", methods=["GET"])
def get_lista_compra(item_id):
    item = db.query(
        "SELECT id, nombre, COALESCE(cantidad, 1) AS cantidad, articulo_id, departamento, COALESCE(estado, 'pendiente') AS estado, COALESCE(encargado, '') AS encargado, creado_at, completado_at FROM lista_compras WHERE id=?",
        [item_id]
    )
    if not item:
        return jsonify({"error": "No encontrado"}), 404
    return jsonify(item[0])


@app.route("/api/lista-compras", methods=["POST"])
def create_lista_compra():
    data = request.get_json(silent=True) or {}
    if not data.get("nombre"):
        return jsonify({"error": "Nombre requerido"}), 400
    cantidad = int(data.get("cantidad", 1))
    departamento = data.get("departamento", "Ferretería")
    encargado = data.get("encargado", "")
    articulo_id = data.get("articulo_id")
    sql = "INSERT INTO lista_compras (nombre, cantidad, departamento, estado, encargado, creado_at, articulo_id) VALUES (?,?,?,?,?,?,?)"
    db.execute(sql, [data["nombre"], cantidad, departamento, "pendiente", encargado, datetime.utcnow().isoformat(), articulo_id])
    return jsonify({"ok": True})


@app.route("/api/lista-compras/<int:item_id>", methods=["PUT"])
def update_lista_compra(item_id):
    data = request.get_json(silent=True) or {}
    estado = data.get("estado")
    nueva_cantidad = data.get("cantidad")
    
    # Obtener el item actual para verificar vinculación con inventario
    res = db.query("SELECT * FROM lista_compras WHERE id=?", [item_id])
    if not res: return jsonify({"error": "No encontrado"}), 404
    item = res[0]

    # Actualizar cantidad si se solicita (desde el botón Editar)
    if nueva_cantidad is not None:
        db.execute("UPDATE lista_compras SET cantidad=? WHERE id=?", [int(nueva_cantidad), item_id])
        item['cantidad'] = int(nueva_cantidad)

    if estado:
        if estado in ['recibido', 'faltante']:
            db.execute("UPDATE lista_compras SET estado=?, completado_at=? WHERE id=?", [estado, datetime.utcnow().isoformat(), item_id])
            
            # Si se marcó como recibido y está vinculado a un artículo, actualizar stock
            if estado == 'recibido' and item.get('articulo_id'):
                art_id = item['articulo_id']
                cant = item['cantidad']
                
                art_res = db.query("SELECT stock FROM articulos WHERE id=?", [art_id])
                if art_res:
                    old_stock = int(art_res[0]['stock'] or 0)
                    new_stock = old_stock + cant
                    
                    db.execute("UPDATE articulos SET stock=? WHERE id=?", [new_stock, art_id])
                    
                    # Registrar movimiento de entrada
                    db.execute("""INSERT INTO movimientos_inventario 
                                  (articulo_id, tipo, cantidad, stock_anterior, stock_nuevo, referencia, motivo)
                                  VALUES (?,?,?,?,?,?,?)""",
                               [art_id, 'entrada', cant, old_stock, new_stock, f"Compra L-C #{item_id}", "Pedido urgente recibido"])
        else:
            db.execute("UPDATE lista_compras SET estado=? WHERE id=?", [estado, item_id])
    return jsonify({"ok": True})


@app.route("/api/lista-compras/<int:item_id>", methods=["DELETE"])
def delete_lista_compra(item_id):
    db.execute("DELETE FROM lista_compras WHERE id=?", [item_id])
    return jsonify({"ok": True})


@app.route("/api/lista-compras/limpiar", methods=["POST"])
def limpiar_lista_compras():
    db.execute("DELETE FROM lista_compras")
    return jsonify({"ok": True})


# ═══════════════════════════════════════
#  CLIENTES Y COTIZACIONES
# ═══════════════════════════════════════

@app.route("/api/clientes", methods=["GET"])
def get_clientes():
    items = db.query("SELECT * FROM clientes ORDER BY nombre")
    return jsonify(items)

@app.route("/api/clientes", methods=["POST"])
def create_cliente():
    data = request.get_json(silent=True) or {}
    db.execute("INSERT INTO clientes (nombre, telefono, rnc_cedula, direccion) VALUES (?,?,?,?)",
               [data.get("nombre"), data.get("telefono"), data.get("rnc_cedula"), data.get("direccion")])
    return jsonify({"ok": True})

@app.route("/api/cotizaciones", methods=["POST"])
def create_cotizacion():
    data = request.get_json(silent=True) or {}
    items_json = json.dumps(data.get("items", []))
    db.execute("INSERT INTO cotizaciones (cliente_id, items, total, validez_dias) VALUES (?,?,?,?)",
               [data.get("cliente_id"), items_json, data.get("total", 0), data.get("validez_dias", 15)])
    return jsonify({"ok": True})

@app.route("/api/cuentas-cobrar", methods=["GET"])
def get_cuentas_cobrar():
    sql = """SELECT cc.*, cl.nombre as cliente_nombre 
             FROM cuentas_cobrar cc 
             LEFT JOIN clientes cl ON cc.cliente_id = cl.id 
             ORDER BY cc.creado_at DESC"""
    items = db.query(sql)
    return jsonify(items)

@app.route("/api/cuentas-cobrar", methods=["POST"])
def create_cuenta_cobrar():
    data = request.get_json(silent=True) or {}
    monto = float(data.get("monto", 0))
    vence = data.get("fecha_vencimiento")
    if not vence:
        vence = (date.today().replace(month=date.today().month+1)).isoformat() # +30 días aprox
        
    db.execute("""INSERT INTO cuentas_cobrar 
                  (cliente_id, concepto, monto, saldo_pendiente, estado, fecha_vencimiento) 
                  VALUES (?,?,?,?,?,?)""",
               [data.get("cliente_id"), data.get("concepto"), monto, monto, "pendiente", vence])
    return jsonify({"ok": True})

@app.route("/api/cuentas-cobrar/<int:cc_id>/pagar", methods=["POST"])
def pagar_cuenta_cobrar(cc_id):
    data = request.get_json(silent=True) or {}
    monto_pago = float(data.get("monto", 0))
    
    res = db.query("SELECT saldo_pendiente FROM cuentas_cobrar WHERE id=?", [cc_id])
    if not res: return jsonify({"error": "No encontrada"}), 404
    
    nuevo_saldo = max(0, res[0]["saldo_pendiente"] - monto_pago)
    estado = "pagada" if nuevo_saldo <= 0 else "pendiente"
    
    db.execute("UPDATE cuentas_cobrar SET saldo_pendiente=?, estado=? WHERE id=?", [nuevo_saldo, estado, cc_id])
    return jsonify({"ok": True, "nuevo_saldo": nuevo_saldo})


# ═══════════════════════════════════════
#  MOVIMIENTOS DE INVENTARIO
# ═══════════════════════════════════════

@app.route("/api/movimientos", methods=["GET"])
def get_movimientos():
    tipo = request.args.get("tipo", "")
    sql = """SELECT m.*, a.nombre as articulo_nombre, a.codigo as articulo_codigo 
             FROM movimientos_inventario m
             LEFT JOIN articulos a ON m.articulo_id = a.id"""
    params = []
    if tipo:
        sql += " WHERE m.tipo = ?"
        params.append(tipo)
    sql += " ORDER BY m.id DESC LIMIT 100"
    items = db.query(sql, params if params else None)
    return jsonify(items)


@app.route("/api/movimientos", methods=["POST"])
def create_movimiento():
    data = request.get_json(silent=True) or {}
    art_id = data.get("articulo_id")
    tipo = data.get("tipo")
    cantidad = int(data.get("cantidad", 0))
    referencia = data.get("referencia", "")
    motivo = data.get("motivo", "")
    if not art_id or not tipo:
        return jsonify({"error": "Datos incompletos"}), 400
    res = db.query("SELECT stock FROM articulos WHERE id=?", [art_id])
    if not res: return jsonify({"error": "Artículo no encontrado"}), 404
    old_stock = int(res[0]['stock'] or 0)
    new_stock = old_stock
    if tipo == 'entrada': new_stock = old_stock + cantidad
    elif tipo == 'salida': new_stock = old_stock - cantidad
    elif tipo == 'ajuste': new_stock = cantidad
    db.execute("UPDATE articulos SET stock=? WHERE id=?", [new_stock, art_id])
    db.execute("""INSERT INTO movimientos_inventario 
                  (articulo_id, tipo, cantidad, stock_anterior, stock_nuevo, referencia, motivo, creado_at)
                  VALUES (?,?,?,?,?,?,?,?)""",
               [art_id, tipo, cantidad, old_stock, new_stock, referencia, motivo, datetime.now().isoformat()])
    return jsonify({"ok": True, "stock_anterior": old_stock, "stock_nuevo": new_stock})


@app.route("/api/facturas/<int:fact_id>", methods=["DELETE"])
def delete_factura(fact_id):
    db.execute("DELETE FROM facturas WHERE id=?", [fact_id])
    return jsonify({"ok": True})


# ═══════════════════════════════════════
#  STATS / DASHBOARD
# ═══════════════════════════════════════

@app.route("/api/stats", methods=["GET"])
def get_stats():
    try:
        arts = db.query("SELECT * FROM articulos")
        provs = db.query("SELECT * FROM proveedores")
        facts_pend = db.query("SELECT COUNT(*) as c FROM facturas WHERE estado='pendiente'")
        
        total_articulos = len(arts)
        total_proveedores = len(provs)
        total_pendientes = facts_pend[0]["c"] if facts_pend else 0
        
        total_costo = sum(float(a.get("precio_costo") or 0) * float(a.get("stock") or 0) for a in arts)
        total_venta = sum(float(a.get("precio_venta") or 0) * float(a.get("stock") or 0) for a in arts)
        margen = total_venta - total_costo
        pct_margen = round((margen / total_costo * 100), 1) if total_costo > 0 else 0
        
        # Dept chart
        depts = ["Repuesto", "Ferretería", "Plomería", "Electricidad", "Hogar"]
        dept_counts = {}
        for d in depts:
            dept_counts[d] = sum(1 for a in arts if a.get("departamento") == d)
        
        # Stock bajo
        bajos = [a for a in arts if a.get("stock", 0) <= a.get("stock_min", 0)]
        
        # Proximos vencimientos
        prox_facts = db.query(
            """SELECT f.*, p.empresa FROM facturas f 
               LEFT JOIN proveedores p ON f.proveedor_id=p.id 
               WHERE f.estado='pendiente' ORDER BY f.fecha_vencimiento ASC LIMIT 6"""
        )
        
        return jsonify({
            "total_articulos": total_articulos,
            "total_proveedores": total_proveedores,
            "total_pendientes": total_pendientes,
            "total_costo": total_costo,
            "total_venta": total_venta,
            "margen": margen,
            "pct_margen": pct_margen,
            "dept_counts": dept_counts,
            "bajos": len(bajos),
            "bajos_lista": [{"nombre": a["nombre"], "stock": a["stock"], "stock_min": a["stock_min"]} for a in bajos[:5]],
            "prox_vencimientos": [
                {
                    "numero": f.get("numero"),
                    "empresa": f.get("empresa"),
                    "monto": f.get("monto"),
                    "fecha_vencimiento": f.get("fecha_vencimiento")
                } for f in prox_facts
            ]
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ═══════════════════════════════════════
#  IMPORT / EXPORT EXCEL
# ═══════════════════════════════════════

# Diccionario para guardar el estado de las importaciones
import_tasks = {}

@app.route("/api/importar-excel", methods=["POST"])
def importar_excel():
    """Import articles from Excel file."""
    if "file" not in request.files:
        return jsonify({"error": "No se envió archivo"}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Archivo vacío"}), 400

    task_id = str(uuid.uuid4())
    file_content = file.read()
    import_tasks[task_id] = {"ok": 0, "err": 0, "total": 0, "status": "procesando", "cancelled": False}

    def process_task(tid, content):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(BytesIO(content))
            ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            import_tasks[tid]["total"] = len(rows)

            for row in rows:
                # Verificar si el usuario canceló la tarea
                if import_tasks[tid].get("cancelled"):
                    import_tasks[tid]["status"] = "cancelado"
                    return

                if not row or len(row) < 2 or not row[1]:
                    import_tasks[tid]["err"] += 1
                    continue
                try:
                    nombre = str(row[1]).strip() if row[1] is not None else ""
                    codigo = str(row[0]).strip() if row[0] is not None else ""
                    dept = str(row[2]).strip() if len(row) > 2 and row[2] is not None else "Ferretería"
                    precio_costo = float(row[3]) if len(row) > 3 and row[3] is not None else 0
                    precio_venta = float(row[4]) if len(row) > 4 and row[4] is not None else 0
                    stock = int(row[5]) if len(row) > 5 and row[5] is not None else 0
                    stock_min = int(row[6]) if len(row) > 6 and row[6] is not None else 5

                    if codigo:
                        existing = db.query("SELECT id FROM articulos WHERE codigo=?", [codigo])
                        if existing:
                            db.execute("UPDATE articulos SET nombre=?, departamento=?, precio_costo=?, precio_venta=?, stock=?, stock_min=? WHERE codigo=?",
                                      [nombre, dept, precio_costo, precio_venta, stock, stock_min, codigo])
                        else:
                            db.execute("INSERT INTO articulos(codigo, nombre, departamento, precio_costo, precio_venta, stock, stock_min) VALUES(?,?,?,?,?,?,?)",
                                      [codigo, nombre, dept, precio_costo, precio_venta, stock, stock_min])
                    else:
                        db.execute("INSERT INTO articulos(nombre, departamento, precio_costo, precio_venta, stock, stock_min) VALUES(?,?,?,?,?,?)",
                                  [nombre, dept, precio_costo, precio_venta, stock, stock_min])
                    import_tasks[tid]["ok"] += 1
                except Exception:
                    import_tasks[tid]["err"] += 1
            import_tasks[tid]["status"] = "completado"
        except Exception as e:
            import_tasks[tid]["status"] = "error"
            import_tasks[tid]["error"] = str(e)

    threading.Thread(target=process_task, args=(task_id, file_content)).start()
    return jsonify({"ok": True, "task_id": task_id})

@app.route("/api/importar-excel/status/<task_id>")
def get_import_status(task_id):
    return jsonify(import_tasks.get(task_id, {"status": "no_encontrado"}))

@app.route("/api/importar-excel/cancel/<task_id>", methods=["POST"])
def cancel_import_task(task_id):
    if task_id in import_tasks:
        import_tasks[task_id]["cancelled"] = True
        return jsonify({"ok": True})
    return jsonify({"error": "Tarea no encontrada"}), 404


@app.route("/api/exportar-articulos", methods=["GET"])
def exportar_articulos():
    """Export all articles as JSON (frontend converts to Excel)."""
    items = db.query(
        """SELECT codigo, nombre, departamento, precio_costo, precio_venta, stock, stock_min, 
                  unidad, descripcion, p.empresa as proveedor 
           FROM articulos a LEFT JOIN proveedores p ON a.proveedor_id=p.id 
           ORDER BY a.nombre"""
    )
    return jsonify(items)


# ═══════════════════════════════════════
#  BACKUPS
# ═══════════════════════════════════════

@app.route("/api/backups", methods=["GET"])
def list_backups():
    """List all available backups."""
    backups = db.query("SELECT id, filename, size, created_at FROM backups ORDER BY created_at DESC")
    return jsonify(backups)


@app.route("/api/backups", methods=["POST"])
def create_backup():
    """Create a full database backup."""
    # Export all data as JSON
    articulos = db.query("SELECT * FROM articulos")
    proveedores = db.query("SELECT * FROM proveedores")
    facturas = db.query("SELECT * FROM facturas")
    
    backup_data = {
        "version": "1.0",
        "created_at": datetime.utcnow().isoformat(),
        "data": {
            "articulos": articulos,
            "proveedores": proveedores,
            "facturas": facturas
        }
    }
    
    json_str = json.dumps(backup_data, ensure_ascii=False, default=str)
    b64_data = base64.b64encode(json_str.encode("utf-8")).decode("utf-8")
    
    filename = f"ferrepro_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    size = len(json_str)
    
    db.execute(
        "INSERT INTO backups (filename, size, created_at, data) VALUES (?,?,?,?)",
        [filename, size, datetime.utcnow().isoformat(), b64_data]
    )
    
    return jsonify({
        "ok": True,
        "filename": filename,
        "size": size,
        "message": "Backup creado exitosamente"
    })


@app.route("/api/backups/<int:backup_id>/restore", methods=["POST"])
def restore_backup(backup_id):
    """Restore database from a backup."""
    bak = db.query("SELECT * FROM backups WHERE id=?", [backup_id])
    if not bak:
        return jsonify({"error": "Backup no encontrado"}), 404
    
    try:
        json_str = base64.b64decode(bak[0]["data"]).decode("utf-8")
        backup_data = json.loads(json_str)
        data = backup_data.get("data", {})
        
        # Clear existing data
        db.execute("DELETE FROM facturas")
        db.execute("DELETE FROM articulos")
        db.execute("DELETE FROM proveedores")
        
        # Restore articulos
        for art in data.get("articulos", []):
            db.execute(
                """INSERT INTO articulos (id, codigo, nombre, departamento, precio_costo, precio_venta, 
                   stock, stock_min, descripcion, proveedor_id, unidad)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                [art.get("id"), art.get("codigo"), art.get("nombre"), art.get("departamento"),
                 art.get("precio_costo"), art.get("precio_venta"), art.get("stock"),
                 art.get("stock_min"), art.get("descripcion"), art.get("proveedor_id"),
                 art.get("unidad", "u.")]
            )
        
        # Restore proveedores
        for prov in data.get("proveedores", []):
            db.execute(
                """INSERT INTO proveedores (id, empresa, contacto, telefono, email, departamento,
                   direccion, rnc, dias_credito, notas)
                   VALUES (?,?,?,?,?,?,?,?,?,?)""",
                [prov.get("id"), prov.get("empresa"), prov.get("contacto"),
                 prov.get("telefono"), prov.get("email"), prov.get("departamento"),
                 prov.get("direccion"), prov.get("rnc"), prov.get("dias_credito"),
                 prov.get("notas")]
            )
        
        # Restore facturas
        for fact in data.get("facturas", []):
            db.execute(
                """INSERT INTO facturas (id, numero, proveedor_id, monto, fecha_emision,
                   fecha_vencimiento, estado, descripcion)
                   VALUES (?,?,?,?,?,?,?,?)""",
                [fact.get("id"), fact.get("numero"), fact.get("proveedor_id"),
                 fact.get("monto"), fact.get("fecha_emision"), fact.get("fecha_vencimiento"),
                 fact.get("estado"), fact.get("descripcion")]
            )
        
        return jsonify({
            "ok": True,
            "message": f"Restaurado: {len(data.get('articulos', []))} artículos, {len(data.get('proveedores', []))} proveedores, {len(data.get('facturas', []))} facturas"
        })
    except Exception as e:
        return jsonify({"error": f"Error al restaurar: {str(e)}"}), 500


@app.route("/api/backups/<int:backup_id>/download", methods=["GET"])
def download_backup(backup_id):
    """Download a backup file."""
    bak = db.query("SELECT * FROM backups WHERE id=?", [backup_id])
    if not bak:
        return jsonify({"error": "Backup no encontrado"}), 404
    
    try:
        json_str = base64.b64decode(bak[0]["data"]).decode("utf-8")
        buffer = BytesIO(json_str.encode("utf-8"))
        return send_file(
            buffer,
            mimetype="application/json",
            as_attachment=True,
            download_name=bak[0]["filename"]
        )
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/backups/restore-upload", methods=["POST"])
def restore_upload_backup():
    """Restore from an uploaded backup file."""
    if "file" not in request.files:
        return jsonify({"error": "No se envió archivo"}), 400
    
    file = request.files["file"]
    try:
        content = file.read().decode("utf-8")
        backup_data = json.loads(content)
        
        if "data" not in backup_data:
            return jsonify({"error": "Formato de backup inválido"}), 400
        
        data = backup_data["data"]
        
        # Clear existing data
        db.execute("DELETE FROM facturas")
        db.execute("DELETE FROM articulos")
        db.execute("DELETE FROM proveedores")
        
        # Restore all
        for art in data.get("articulos", []):
            db.execute(
                "INSERT INTO articulos (id, codigo, nombre, departamento, precio_costo, precio_venta, stock, stock_min, descripcion, proveedor_id, unidad) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                [art.get("id"), art.get("codigo"), art.get("nombre"), art.get("departamento"),
                 art.get("precio_costo"), art.get("precio_venta"), art.get("stock"),
                 art.get("stock_min"), art.get("descripcion"), art.get("proveedor_id"),
                 art.get("unidad", "u.")]
            )
        
        for prov in data.get("proveedores", []):
            db.execute(
                "INSERT INTO proveedores (id, empresa, contacto, telefono, email, departamento, direccion, rnc, dias_credito, notas) VALUES (?,?,?,?,?,?,?,?,?,?)",
                [prov.get("id"), prov.get("empresa"), prov.get("contacto"),
                 prov.get("telefono"), prov.get("email"), prov.get("departamento"),
                 prov.get("direccion"), prov.get("rnc"), prov.get("dias_credito"),
                 prov.get("notas")]
            )
        
        for fact in data.get("facturas", []):
            db.execute(
                "INSERT INTO facturas (id, numero, proveedor_id, monto, fecha_emision, fecha_vencimiento, estado, descripcion) VALUES (?,?,?,?,?,?,?,?)",
                [fact.get("id"), fact.get("numero"), fact.get("proveedor_id"),
                 fact.get("monto"), fact.get("fecha_emision"), fact.get("fecha_vencimiento"),
                 fact.get("estado"), fact.get("descripcion")]
            )
        
        return jsonify({"ok": True, "message": "Backup restaurado exitosamente"})
    except json.JSONDecodeError:
        return jsonify({"error": "El archivo no es un JSON válido"}), 400
    except Exception as e:
        return jsonify({"error": f"Error al restaurar: {str(e)}"}), 500


@app.route("/api/backups/<int:backup_id>", methods=["DELETE"])
def delete_backup(backup_id):
    """Delete a backup record (not the database data)."""
    db.execute("DELETE FROM backups WHERE id=?", [backup_id])
    return jsonify({"ok": True})


# ═══════════════════════════════════════
#  CAMBIAR CONTRASEÑA
# ═══════════════════════════════════════

@app.route("/api/cambiar-pass", methods=["POST"])
def cambiar_password():
    data = request.get_json(silent=True) or {}
    if data.get("password", ""):
        return jsonify({"ok": True})
    return jsonify({"ok": False, "error": "Contraseña requerida"}), 400


# ═══════════════════════════════════════
#  HEALTH CHECK
# ═══════════════════════════════════════

@app.route("/api/health", methods=["GET"])
def health():
    try:
        tables = db.init_db()
        return jsonify({
                "status": "ok",
                "tables": [t["name"] for t in tables] if tables else [],
                "db_url": db.SUPABASE_DATABASE_URL[:20] + "..." if db.SUPABASE_DATABASE_URL else "not configured"
            })
    except Exception as e:
        return jsonify({"status": "error", "error": str(e)}), 500


@app.route("/")
def serve_frontend():
    return render_template('index.html')


# ─── Init DB on startup ───
def initialize():
    try:
        tables = db.init_db()
        print(f"[OK] DB connected. Tables: {[t['name'] for t in tables] if tables else 'none'}")
    except Exception as e:
        print(f"[ERROR] DB init error: {e}")

# Inicializar base de datos al importar (necesario para Vercel)
initialize()


if __name__ == "__main__":
    # Ejecución local
    port = int(os.getenv("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=True)
